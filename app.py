from flask import Flask, render_template, request, redirect, url_for, session, flash, send_from_directory
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
import sqlite3
from datetime import datetime
import os
import random
import re

try:
    import PyPDF2
except Exception:
    PyPDF2 = None

try:
    from docx import Document
except Exception:
    Document = None

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, 'database.db')
UPLOAD_FOLDER = os.path.join(BASE_DIR, 'uploads')
ALLOWED_EXTENSIONS = {'txt', 'pdf', 'docx'}
ADMIN_EMAIL = 'ns.saridar@gmail.com'
SECRET_KEY = 'smc-v8-dev-secret-change-this'

app = Flask(__name__)
app.config['SECRET_KEY'] = SECRET_KEY
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
app.config['MAX_CONTENT_LENGTH'] = 20 * 1024 * 1024

os.makedirs(UPLOAD_FOLDER, exist_ok=True)


def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS


def now():
    return datetime.now().strftime('%Y-%m-%d %H:%M:%S')


def init_db():
    conn = get_db()
    cur = conn.cursor()

    cur.execute('''
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            email TEXT UNIQUE NOT NULL,
            password_hash TEXT NOT NULL,
            plan TEXT NOT NULL DEFAULT 'free',
            role TEXT NOT NULL DEFAULT 'user',
            is_active INTEGER NOT NULL DEFAULT 1,
            created_at TEXT NOT NULL,
            usage_quiz INTEGER NOT NULL DEFAULT 0,
            usage_crisis INTEGER NOT NULL DEFAULT 0,
            usage_debate INTEGER NOT NULL DEFAULT 0,
            usage_speech INTEGER NOT NULL DEFAULT 0
        )
    ''')

    cur.execute('''
        CREATE TABLE IF NOT EXISTS conferences (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            name TEXT NOT NULL,
            committee TEXT NOT NULL,
            country TEXT NOT NULL,
            topic TEXT NOT NULL,
            event_datetime TEXT NOT NULL,
            summary TEXT,
            created_at TEXT NOT NULL,
            FOREIGN KEY(user_id) REFERENCES users(id)
        )
    ''')

    cur.execute('''
        CREATE TABLE IF NOT EXISTS library_files (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            original_name TEXT NOT NULL,
            stored_name TEXT NOT NULL,
            file_type TEXT NOT NULL,
            extracted_text TEXT,
            uploaded_at TEXT NOT NULL,
            FOREIGN KEY(user_id) REFERENCES users(id)
        )
    ''')

    cur.execute('''
        CREATE TABLE IF NOT EXISTS activity_log (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER,
            action TEXT NOT NULL,
            details TEXT,
            created_at TEXT NOT NULL,
            FOREIGN KEY(user_id) REFERENCES users(id)
        )
    ''')

    conn.commit()

    admin = conn.execute('SELECT * FROM users WHERE email = ?', (ADMIN_EMAIL,)).fetchone()
    if not admin:
        conn.execute(
            'INSERT INTO users (name, email, password_hash, plan, role, is_active, created_at) VALUES (?, ?, ?, ?, ?, ?, ?)',
            ('Nicolas Saridar', ADMIN_EMAIL, generate_password_hash('admin12345'), 'premium', 'admin', 1, now())
        )
        conn.commit()

    conn.close()


def current_user():
    uid = session.get('user_id')
    if not uid:
        return None
    conn = get_db()
    user = conn.execute('SELECT * FROM users WHERE id = ?', (uid,)).fetchone()
    conn.close()
    return user


def login_required(fn):
    from functools import wraps
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if not session.get('user_id'):
            flash('Please log in first.', 'error')
            return redirect(url_for('login'))
        return fn(*args, **kwargs)
    return wrapper


def admin_required(fn):
    from functools import wraps
    @wraps(fn)
    def wrapper(*args, **kwargs):
        user = current_user()
        if not user or user['role'] != 'admin':
            flash('Admin access only.', 'error')
            return redirect(url_for('dashboard'))
        return fn(*args, **kwargs)
    return wrapper


def log_action(user_id, action, details=''):
    conn = get_db()
    conn.execute(
        'INSERT INTO activity_log (user_id, action, details, created_at) VALUES (?, ?, ?, ?)',
        (user_id, action, details, now())
    )
    conn.commit()
    conn.close()


def increment_usage(user_id, column):
    conn = get_db()
    conn.execute(f'UPDATE users SET {column} = {column} + 1 WHERE id = ?', (user_id,))
    conn.commit()
    conn.close()


def premium_allowed(user, feature):
    if user['role'] == 'admin' or user['plan'] == 'premium':
        return True, ''

    limits = {
        'quiz': ('usage_quiz', 3),
        'crisis': ('usage_crisis', 2),
        'debate': ('usage_debate', 2),
        'speech': ('usage_speech', 3),
    }
    col, limit = limits[feature]
    if user[col] >= limit:
        return False, f'Free plan limit reached for {feature}. Ask admin for premium.'
    return True, ''


def extract_text(file_path, ext):
    ext = ext.lower()
    try:
        if ext == 'txt':
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()[:50000]
        if ext == 'pdf' and PyPDF2:
            text = []
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages[:20]:
                    text.append(page.extract_text() or '')
            return '\n'.join(text)[:50000]
        if ext == 'docx' and Document:
            doc = Document(file_path)
            return '\n'.join(p.text for p in doc.paragraphs)[:50000]
    except Exception:
        return ''
    return ''


def build_context_text(user_id, file_ids=None):
    conn = get_db()
    if file_ids:
        placeholders = ','.join('?' for _ in file_ids)
        rows = conn.execute(
            f'SELECT extracted_text, original_name FROM library_files WHERE user_id = ? AND id IN ({placeholders}) ORDER BY id DESC',
            [user_id] + file_ids
        ).fetchall()
    else:
        rows = conn.execute(
            'SELECT extracted_text, original_name FROM library_files WHERE user_id = ? ORDER BY id DESC LIMIT 5',
            (user_id,)
        ).fetchall()
    conn.close()
    chunks = []
    for row in rows:
        if row['extracted_text']:
            chunks.append(f"Source File: {row['original_name']}\n{row['extracted_text'][:6000]}")
    return '\n\n'.join(chunks)[:18000]


def sentence_split(text):
    bits = re.split(r'(?<=[.!?])\s+|\n+', text.strip())
    return [b.strip() for b in bits if len(b.strip()) > 15]


def looks_like_junk(line):
    s = re.sub(r'\s+', ' ', line.strip())
    if not s:
        return True
    lower = s.lower()
    junk_fragments = [
        'source file:', '.docx', '.pdf', '.txt', 'page ', 'part ', 'chapter ',
        'copyright', 'all rights reserved', 'table of contents'
    ]
    if any(frag in lower for frag in junk_fragments):
        return True
    letters = sum(ch.isalpha() for ch in s)
    if letters < 12:
        return True
    if sum(ch.isupper() for ch in s if ch.isalpha()) > max(10, letters * 0.75):
        return True
    if len(s.split()) < 6:
        return True
    return False


def clean_context(text):
    pieces = []
    seen = set()
    for raw in sentence_split(text):
        line = re.sub(r'\s+', ' ', raw).strip(' -•\t')
        if looks_like_junk(line):
            continue
        line = re.sub(r'^(source|document|file)\s*:\s*', '', line, flags=re.I)
        line = re.sub(r'\s+', ' ', line).strip()
        if len(line) < 40 or len(line) > 360:
            continue
        key = line.lower()
        if key in seen:
            continue
        seen.add(key)
        pieces.append(line)
    return pieces


KEYWORDS = {
    'action': ['must', 'should', 'propose', 'recommend', 'call for', 'urge', 'support', 'establish', 'create', 'adopt', 'implement', 'strengthen', 'coordinate', 'protect', 'respond'],
    'risk': ['threat', 'risk', 'danger', 'harm', 'attack', 'escalation', 'crisis', 'disruption', 'instability', 'vulnerability'],
    'stance': ['supports', 'opposes', 'believes', 'argues', 'maintains', 'rejects', 'favours', 'prefers', 'insists'],
    'law': ['article', 'charter', 'law', 'legal', 'international humanitarian law', 'sovereignty', 'jurisdiction'],
    'coop': ['allies', 'coordination', 'multilateral', 'nato', 'united nations', 'committee', 'partners', 'collective']
}


def sentence_score(sentence, topic=''):
    lower = sentence.lower()
    score = 0
    for words in KEYWORDS.values():
        for w in words:
            if w in lower:
                score += 2
    if topic:
        for tok in re.findall(r'[a-zA-Z]{4,}', topic.lower()):
            if tok in lower:
                score += 3
    if any(ch.isdigit() for ch in sentence):
        score += 1
    score += min(len(sentence.split()) // 8, 3)
    return score


def classify_sentence(sentence):
    lower = sentence.lower()
    if any(w in lower for w in KEYWORDS['action']):
        return 'policy'
    if any(w in lower for w in KEYWORDS['risk']):
        return 'risk'
    if any(w in lower for w in KEYWORDS['law']):
        return 'law'
    if any(w in lower for w in KEYWORDS['stance']):
        return 'stance'
    return 'fact'


def compress_sentence(sentence, max_words=16):
    cleaned = re.sub(r'\s+', ' ', sentence).strip()
    cleaned = re.sub(r'^(Spain|The delegation|The committee|This committee)\s+', '', cleaned, flags=re.I)
    words = cleaned.split()
    if len(words) <= max_words:
        return cleaned.rstrip('. ')
    return ' '.join(words[:max_words]).rstrip(',;:. ') + '...'



def normalize_phrase(text, max_words=16):
    text = re.sub(r'\s+', ' ', text).strip(' .,:;-"\'')
    words = text.split()
    if len(words) > max_words:
        text = ' '.join(words[:max_words]).rstrip(' ,;:') + '…'
    return text

def titleish(text):
    return text[:1].upper() + text[1:] if text else text

def unique_keep_order(items, key=lambda x: x):
    seen = set()
    out = []
    for item in items:
        k = key(item)
        if not k or k in seen:
            continue
        seen.add(k)
        out.append(item)
    return out

def extract_structured_points(sentences, topic=''):
    points = []
    topic_terms = {w.lower() for w in re.findall(r'[A-Za-z]{4,}', topic)}
    for sent in sentences:
        lower = sent.lower()
        score = sentence_score(sent, topic)
        if topic_terms:
            score += sum(1 for t in topic_terms if t in lower)
        kind = classify_sentence(sent)
        action_match = re.search(r'\b(should|must|propose|recommend|urge|support|create|adopt|implement|strengthen|coordinate|protect|establish|develop|increase|improve)\b(.+)', sent, flags=re.I)
        if action_match:
            action = normalize_phrase(action_match.group(0))
            points.append({'kind': 'policy', 'text': sent, 'summary': titleish(action), 'score': score + 3})
        risk_match = re.search(r'\b(risk|threat|danger|weakness|challenge|pressure|vulnerability|instability|attack|crisis|failure)\b(.+)', sent, flags=re.I)
        if risk_match:
            risk = normalize_phrase(risk_match.group(0))
            points.append({'kind': 'risk', 'text': sent, 'summary': titleish(risk), 'score': score + 2})
        stance_match = re.search(r'\b(spain|delegation|country|state|government|nato|committee)\b(.+)', sent, flags=re.I)
        if stance_match or kind in ('stance', 'law', 'fact'):
            summ = compress_sentence(sent, 16)
            points.append({'kind': kind, 'text': sent, 'summary': titleish(summ), 'score': score + 1})
    if not points:
        for sent in sentences:
            points.append({'kind': classify_sentence(sent), 'text': sent, 'summary': titleish(compress_sentence(sent, 16)), 'score': sentence_score(sent, topic)})
    points = sorted(points, key=lambda p: (p['score'], len(p['text'])), reverse=True)
    return unique_keep_order(points, key=lambda p: (p['kind'], p['summary'].lower()))

def build_argument_bank(topic, context, country=''):
    sentences = clean_context(context)
    structured = extract_structured_points(sentences, topic)
    top = structured[:14]
    policy = [p for p in top if p['kind'] == 'policy'][:5]
    risks = [p for p in top if p['kind'] == 'risk'][:4]
    law = [p for p in top if p['kind'] == 'law'][:3]
    stance = [p for p in top if p['kind'] in ('stance', 'fact', 'law')][:4]
    facts = [p for p in top if p['kind'] == 'fact'][:4]
    return {
        'all': top,
        'policy': policy,
        'risks': risks,
        'law': law,
        'stance': stance,
        'facts': facts,
        'summary': build_research_profile(topic, context, country).get('summary')
    }

def build_research_profile(topic, context, country=''):
    sentences = clean_context(context)
    structured = extract_structured_points(sentences, topic)
    top = structured[:12]
    if not top:
        return {
            'insights': [],
            'summary': f'{country or "The delegation"} needs a structured, evidence-based position on {topic}.',
            'actions': [
                f'Clarify your delegation’s position on {topic}.',
                'Support every claim with evidence from your research.',
                'Prepare one practical short-term response and one long-term solution.'
            ]
        }

    policy = [i for i in top if i['kind'] == 'policy'][:4]
    risks = [i for i in top if i['kind'] == 'risk'][:3]
    stance = [i for i in top if i['kind'] in ('stance', 'law', 'fact')][:3]

    summary_parts = []
    if stance:
        summary_parts.append(f"Key position: {stance[0]['summary']}")
    if policy:
        summary_parts.append(f"Preferred move: {policy[0]['summary']}")
    if risks:
        summary_parts.append(f"Main pressure point: {risks[0]['summary']}")

    actions = [i['summary'] for i in unique_keep_order(policy + stance + risks, key=lambda x: x['summary'].lower())[:5]]
    return {
        'insights': top,
        'summary': ' | '.join(summary_parts) if summary_parts else compress_sentence(top[0]['text'], 18),
        'actions': actions or [compress_sentence(top[0]['text'], 14)]
    }


def make_distractors(correct, insight_pool):
    pool = []
    for item in insight_pool:
        cand = normalize_phrase(item['summary'])
        if cand.lower() != correct.lower() and cand not in pool and len(cand.split()) >= 4:
            pool.append(cand)
    generic = [
        'Escalate first and justify the move later',
        'Delay all action until absolute certainty appears',
        'Rely on slogans rather than a concrete policy step',
        'Ignore allies, procedure, and implementation details',
        'Treat the issue as rhetorical rather than operational'
    ]
    for g in generic:
        if g.lower() != correct.lower() and g not in pool:
            pool.append(g)
    return pool[:3]


def improved_quiz(topic, context):
    bank = build_argument_bank(topic, context)
    insights = bank['all']
    if not insights:
        insights = [
            {'kind': 'policy', 'text': f'Delegates should build a clear, evidence-based strategy on {topic}.', 'summary': f'Build a clear evidence-based strategy on {topic}'},
            {'kind': 'risk', 'text': f'The biggest weakness on {topic} is speaking without evidence or structure.', 'summary': f'Avoid unsupported and unstructured claims on {topic}'},
            {'kind': 'stance', 'text': f'Effective delegates connect national interest to realistic multilateral action on {topic}.', 'summary': 'Link national interest to realistic multilateral action'}
        ]

    questions = []
    used = set()

    def add_question(kind_label, question, statement, correct, pool, explanation):
        key = (question, correct)
        if key in used or len(correct.split()) < 3:
            return
        options = [correct] + make_distractors(correct, pool)
        options = unique_keep_order(options, key=lambda x: x.lower())
        if len(options) < 4:
            return
        options = options[:4]
        random.shuffle(options)
        questions.append({
            'question': question,
            'statement': statement,
            'options': options,
            'answer': correct,
            'explanation': explanation,
            'kind': kind_label
        })
        used.add(key)

    for item in bank['policy'][:2]:
        add_question(
            'Policy',
            'Based on the research, which move is the strongest first step for the delegation?',
            item['text'],
            normalize_phrase(item['summary'], 18),
            insights,
            f"This is best supported because the research explicitly points toward this policy direction: {item['text']}"
        )

    for item in bank['risks'][:1]:
        add_question(
            'Risk',
            'Which vulnerability or pressure point does the research most clearly warn the delegate about?',
            item['text'],
            normalize_phrase(item['summary'], 18),
            insights,
            f"The research frames this as a central pressure point: {item['text']}"
        )

    for item in bank['law'][:1] or bank['stance'][:1]:
        add_question(
            'Principle',
            'Which principle or position is most consistent with the research?',
            item['text'],
            normalize_phrase(item['summary'], 18),
            insights,
            f"This option best matches the position developed in the source material: {item['text']}"
        )

    for item in bank['facts'][:1] or insights[:1]:
        add_question(
            'Interpretation',
            'Which conclusion would a strong delegate most reasonably draw from this research?',
            item['text'],
            normalize_phrase(item['summary'], 18),
            insights,
            f"This is the most defensible conclusion because it directly reflects the argument in the research: {item['text']}"
        )

    return questions[:5]


def improved_crisis(topic, country, committee, context):
    bank = build_argument_bank(topic, context, country)
    profile = build_research_profile(topic, context, country)
    policy = bank['policy'][:3]
    risks = bank['risks'][:3]
    stance = bank['stance'][:2]

    topic_lower = (topic or '').lower()
    committee_upper = (committee or '').upper()
    if 'nato' in committee_upper or 'article 5' in topic_lower or 'cyber' in topic_lower:
        attacker = 'the United States' if country.lower() == 'spain' else 'a NATO member state'
        trigger = risks[0]['summary'] if risks else 'forensic teams confirm coordinated access to military and civilian systems'
        first_move = policy[0]['summary'] if policy else 'demand immediate allied consultation, evidence review, and proportional response planning'
        scenario = (
            f"Emergency brief for {committee}: at 04:20 GMT, malware seeded through a trusted supply-chain update breaches {country}'s defence network, air-traffic coordination layer, and portions of its national energy dispatch system. "
            f"Within forty minutes, leaked forensic markers lead several capitals to accuse {attacker}, while other delegations warn that a premature accusation could fracture alliance unity before attribution is fully verified. "
            f"Spanish officials privately argue that the pattern is too coordinated to dismiss as criminal sabotage, and multiple ambassadors are now openly asking whether the incident crosses the political threshold for Article 5 consultation. "
            f"{trigger}. Markets are reacting, media outlets are broadcasting the breach as a test of allied credibility, and your delegation has ten minutes before closed consultations begin. One weak phrase could isolate {country}; one reckless phrase could push the committee toward a response it cannot politically reverse."
        )
        intel = (
            f"Research profile: {profile['summary']}. Inside the room, delegates are split into three camps: one bloc wants public attribution and deterrent language immediately, another wants a technical review before any alliance trigger is discussed, and a third is trying to keep the crisis below the threshold of formal collective-defence commitments."
        )
        tasks = [
            f"State {country}'s first official line in a way that sounds decisive but does not overclaim attribution.",
            f"Propose the immediate allied mechanism or committee step that should happen first: {first_move}.",
            "Explain whether you support Article 5 language now, later, or not at all — and justify the threshold.",
            "Name the parallel diplomatic and operational measures that should happen in the first response window."
        ]
    else:
        trigger = risks[0]['summary'] if risks else f'a rapidly escalating development linked to {topic}'
        first_move = policy[0]['summary'] if policy else f'present a realistic first response on {topic}'
        framing = stance[0]['summary'] if stance else f'{country} must sound credible, measured, and useful'
        scenario = (
            f"Emergency brief for {committee}: in the last hour, {trigger.lower()} has moved the committee from general principle to immediate decision-making. "
            f"Multiple delegations now want to know exactly what {country} will do first, what line it will not cross, and whether its position can survive pressure from both hawks and skeptics. "
            f"{framing}. You have one intervention to shape the room before informal consultations harden into rival blocs."
        )
        intel = f"Research profile: {profile['summary']}. The room is no longer rewarding abstract values; it is judging whether {country} can turn principle into implementation under pressure."
        tasks = [
            "Open with one sentence that sounds diplomatic but decisive.",
            f"Propose one immediate committee step grounded in your research: {first_move}.",
            "Defend your response against criticism that it is either too weak or too escalatory.",
            "State what must be verified before the committee locks itself into irreversible language."
        ]

    benchmarks = unique_keep_order(
        [normalize_phrase((stance[0]['summary'] if stance else profile['summary']), 16)] +
        [normalize_phrase(p['summary'], 16) for p in policy] +
        [normalize_phrase(r['summary'], 16) for r in risks],
        key=lambda x: x.lower()
    )[:5]

    return {
        'scenario': scenario,
        'intel': intel,
        'tasks': tasks,
        'benchmarks': benchmarks
    }


def improved_debate(topic, your_country, opponent_country, context):
    bank = build_argument_bank(topic, context, your_country)
    profile = build_research_profile(topic, context, your_country)
    strongest = [normalize_phrase(x['summary'], 16) for x in bank['policy'][:3] or bank['stance'][:2]]
    pressure = [normalize_phrase(x['summary'], 16) for x in bank['risks'][:3]]
    anchor = strongest[0] if strongest else f'produce a clearer and more realistic policy on {topic}'
    vulnerability = pressure[0] if pressure else f'explain how its plan avoids weak enforcement and vague implementation'

    topic_lower = (topic or '').lower()
    if 'nato' in topic_lower or 'article 5' in topic_lower or 'cyber' in topic_lower:
        opener = (
            f"Delegate of {opponent_country}: {your_country} asks this committee to trust a doctrine of discipline, proportionality, and alliance unity. "
            f"But when the pressure becomes real — when civilian systems are down, attribution is strong but not yet perfect, and allies are already dividing — your framework still hesitates at the exact point where deterrence is supposed to matter. "
            f"Your own research leans toward {anchor.lower()}, yet it also leaves one dangerous opening: how do you answer {vulnerability.lower()} without either paralyzing the alliance or pushing it into escalation without consensus? "
            f"So answer clearly: what threshold would make {your_country} support collective action, what exact first step would it endorse inside NATO, and why should this committee follow your caution instead of a harder and faster line?"
        )
        pressure_points = unique_keep_order(
            pressure + [
                f"Force {your_country} to define its Article 5 threshold instead of hiding behind general principles.",
                'Attack any gap between attribution standards and the speed of the response.',
                'Question whether caution becomes paralysis when deterrence must be visible.',
                'Push the delegate to name one implementable allied action, not just a value statement.'
            ],
            key=lambda x: x.lower()
        )[:5]
        counters = unique_keep_order((strongest or []) + [
            'Define a threshold that links evidence, proportionality, and alliance consultation rather than pure delay.',
            'Argue that disciplined attribution protects alliance unity and makes any collective move more credible.',
            'Name one immediate NATO mechanism or consultation step to prove your plan is operational, not rhetorical.'
        ] + ([normalize_phrase(profile['summary'], 18)] if profile.get('summary') else []), key=lambda x: x.lower())[:5]
    else:
        opener = (
            f"Delegate of {opponent_country}: We respect {your_country}, but this committee still lacks proof that your approach on {topic} is the most workable one. "
            f"Your own research points toward {anchor.lower()}, yet it also leaves you exposed on one central question: how do you address {vulnerability.lower()} without creating new instability? "
            f"If your model is truly credible, explain the first implementable step and why the committee should trust it over a harder-line alternative."
        )
        pressure_points = unique_keep_order(
            pressure + [
                f'Question whether {your_country} can translate principle into implementation.',
                'Challenge enforcement, timing, attribution, or coordination gaps.',
                'Push the delegate to prove realism, not just rhetoric.'
            ],
            key=lambda x: x.lower()
        )[:5]
        counters = strongest or [
            'Anchor your answer in one immediate and realistic policy step.',
            'Show how your model lowers risk without reckless escalation.'
        ]
        if profile.get('summary'):
            counters = unique_keep_order(counters + [normalize_phrase(profile['summary'], 18)], key=lambda x: x.lower())
        counters = counters[:5]

    return {
        'opener': opener,
        'pressure_points': pressure_points,
        'best_counters': counters
    }


def fallback_speech_feedback(text):
    words = len(re.findall(r'\w+', text))
    paragraphs = [p.strip() for p in text.split('\n') if p.strip()]
    score = min(95, max(60, 62 + len(paragraphs) * 4 + min(words, 400) // 20))
    strengths = [
        'Your speech has a clear voice and a visible diplomatic tone.',
        'There is an identifiable structure rather than random points.',
        'Several lines sound persuasive and presentation-ready.'
    ]
    improvements = [
        'Add one sharper hook in the opening sentence.',
        'Use more country-specific policy details to sound authoritative.',
        'End with a stronger final line that sounds quotable.'
    ]
    revised_opening = 'Honourable Chair, distinguished delegates, when instability reaches our screens it feels distant, but for millions it is daily life, and this committee cannot afford the comfort of delay.'
    return {
        'score': score,
        'strengths': strengths,
        'improvements': improvements,
        'revised_opening': revised_opening
    }


@app.context_processor
def inject_globals():
    return {'current_user': current_user()}


@app.route('/')
def index():
    if session.get('user_id'):
        return redirect(url_for('dashboard'))
    return render_template('landing.html')


@app.route('/signup', methods=['GET', 'POST'])
def signup():
    if request.method == 'POST':
        name = request.form.get('name', '').strip()
        email = request.form.get('email', '').strip().lower()
        password = request.form.get('password', '')

        if not name or not email or not password:
            flash('Please fill every field.', 'error')
            return redirect(url_for('signup'))
        if len(password) < 6:
            flash('Password must be at least 6 characters.', 'error')
            return redirect(url_for('signup'))

        conn = get_db()
        existing = conn.execute('SELECT id FROM users WHERE email = ?', (email,)).fetchone()
        if existing:
            conn.close()
            flash('That email already exists.', 'error')
            return redirect(url_for('signup'))

        role = 'admin' if email == ADMIN_EMAIL else 'user'
        plan = 'premium' if email == ADMIN_EMAIL else 'free'
        conn.execute(
            'INSERT INTO users (name, email, password_hash, plan, role, created_at) VALUES (?, ?, ?, ?, ?, ?)',
            (name, email, generate_password_hash(password), plan, role, now())
        )
        conn.commit()
        user = conn.execute('SELECT * FROM users WHERE email = ?', (email,)).fetchone()
        conn.close()
        session['user_id'] = user['id']
        log_action(user['id'], 'signup', f'{email} created account')
        flash('Account created successfully.', 'success')
        return redirect(url_for('dashboard'))
    return render_template('auth.html', mode='signup')


@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        email = request.form.get('email', '').strip().lower()
        password = request.form.get('password', '')
        conn = get_db()
        user = conn.execute('SELECT * FROM users WHERE email = ?', (email,)).fetchone()
        conn.close()
        if not user or not check_password_hash(user['password_hash'], password):
            flash('Invalid email or password.', 'error')
            return redirect(url_for('login'))
        if not user['is_active']:
            flash('This account is suspended.', 'error')
            return redirect(url_for('login'))
        session['user_id'] = user['id']
        log_action(user['id'], 'login', f'{email} logged in')
        flash('Welcome back.', 'success')
        if user['role'] == 'admin':
            return redirect(url_for('admin_panel'))
        return redirect(url_for('dashboard'))
    return render_template('auth.html', mode='login')


@app.route('/logout')
def logout():
    uid = session.get('user_id')
    session.clear()
    if uid:
        log_action(uid, 'logout', 'User logged out')
    flash('Logged out successfully.', 'success')
    return redirect(url_for('index'))


@app.route('/dashboard')
@login_required
def dashboard():
    user = current_user()
    conn = get_db()
    conferences = conn.execute('SELECT * FROM conferences WHERE user_id = ? ORDER BY created_at DESC', (user['id'],)).fetchall()
    files = conn.execute('SELECT * FROM library_files WHERE user_id = ? ORDER BY uploaded_at DESC', (user['id'],)).fetchall()
    stats = {
        'conference_count': conn.execute('SELECT COUNT(*) AS c FROM conferences WHERE user_id = ?', (user['id'],)).fetchone()['c'],
        'file_count': conn.execute('SELECT COUNT(*) AS c FROM library_files WHERE user_id = ?', (user['id'],)).fetchone()['c'],
    }
    conn.close()
    return render_template('dashboard.html', conferences=conferences, files=files, stats=stats)


@app.route('/conference/create', methods=['POST'])
@login_required
def create_conference():
    user = current_user()
    name = request.form.get('name', '').strip()
    committee = request.form.get('committee', '').strip()
    country = request.form.get('country', '').strip()
    topic = request.form.get('topic', '').strip()
    event_datetime = request.form.get('event_datetime', '').strip()
    summary = request.form.get('summary', '').strip()

    if not all([name, committee, country, topic, event_datetime]):
        flash('Please fill all conference fields including date and time.', 'error')
        return redirect(url_for('dashboard'))

    try:
        dt = datetime.strptime(event_datetime, '%Y-%m-%dT%H:%M')
        event_value = dt.strftime('%Y-%m-%d %H:%M:%S')
    except ValueError:
        flash('Invalid date and time format.', 'error')
        return redirect(url_for('dashboard'))

    conn = get_db()
    conn.execute(
        'INSERT INTO conferences (user_id, name, committee, country, topic, event_datetime, summary, created_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?)',
        (user['id'], name, committee, country, topic, event_value, summary, now())
    )
    conn.commit()
    conn.close()
    log_action(user['id'], 'create_conference', f'{name} / {committee}')
    flash('Conference created successfully.', 'success')
    return redirect(url_for('dashboard'))


@app.route('/library/upload', methods=['POST'])
@login_required
def upload_library():
    user = current_user()
    files = request.files.getlist('files')
    if not files or files == [None]:
        flash('Please select at least one file.', 'error')
        return redirect(url_for('dashboard'))

    uploaded = 0
    conn = get_db()
    for file in files:
        if not file or not file.filename:
            continue
        if not allowed_file(file.filename):
            continue
        original_name = file.filename
        ext = original_name.rsplit('.', 1)[1].lower()
        stored_name = f"{int(datetime.now().timestamp()*1000)}_{secure_filename(original_name)}"
        path = os.path.join(app.config['UPLOAD_FOLDER'], stored_name)
        file.save(path)
        extracted = extract_text(path, ext)
        conn.execute(
            'INSERT INTO library_files (user_id, original_name, stored_name, file_type, extracted_text, uploaded_at) VALUES (?, ?, ?, ?, ?, ?)',
            (user['id'], original_name, stored_name, ext, extracted, now())
        )
        uploaded += 1
    conn.commit()
    conn.close()
    log_action(user['id'], 'upload_files', f'Uploaded {uploaded} file(s)')
    flash(f'Uploaded {uploaded} file(s).', 'success' if uploaded else 'error')
    return redirect(url_for('dashboard'))


@app.route('/library/download/<int:file_id>')
@login_required
def download_file(file_id):
    user = current_user()
    conn = get_db()
    row = conn.execute('SELECT * FROM library_files WHERE id = ? AND user_id = ?', (file_id, user['id'])).fetchone()
    conn.close()
    if not row:
        flash('File not found.', 'error')
        return redirect(url_for('dashboard'))
    return send_from_directory(app.config['UPLOAD_FOLDER'], row['stored_name'], as_attachment=True, download_name=row['original_name'])


@app.route('/speech-lab', methods=['GET', 'POST'])
@login_required
def speech_lab():
    result = None
    user = current_user()
    if request.method == 'POST':
        allowed, msg = premium_allowed(user, 'speech')
        if not allowed:
            flash(msg, 'error')
            return redirect(url_for('speech_lab'))
        text = request.form.get('speech_text', '').strip()
        if not text:
            flash('Paste a speech first.', 'error')
            return redirect(url_for('speech_lab'))
        result = fallback_speech_feedback(text)
        increment_usage(user['id'], 'usage_speech')
        log_action(user['id'], 'speech_lab', 'Analyzed a speech')
    return render_template('speech_lab.html', result=result)


@app.route('/quiz', methods=['GET', 'POST'])
@login_required
def quiz():
    user = current_user()
    conn = get_db()
    files = conn.execute('SELECT id, original_name FROM library_files WHERE user_id = ? ORDER BY uploaded_at DESC', (user['id'],)).fetchall()
    conn.close()
    quiz_data = None
    topic = ''
    if request.method == 'POST':
        allowed, msg = premium_allowed(user, 'quiz')
        if not allowed:
            flash(msg, 'error')
            return redirect(url_for('quiz'))
        topic = request.form.get('topic', '').strip() or 'Your MUN topic'
        raw_ids = request.form.getlist('file_ids')
        file_ids = [int(x) for x in raw_ids if x.isdigit()]
        context = build_context_text(user['id'], file_ids)
        quiz_data = improved_quiz(topic, context)
        increment_usage(user['id'], 'usage_quiz')
        log_action(user['id'], 'quiz', f'Generated quiz on {topic}')
    return render_template('quiz.html', files=files, quiz_data=quiz_data, topic=topic)


@app.route('/crisis', methods=['GET', 'POST'])
@login_required
def crisis():
    user = current_user()
    result = None
    conn = get_db()
    files = conn.execute('SELECT id, original_name FROM library_files WHERE user_id = ? ORDER BY uploaded_at DESC', (user['id'],)).fetchall()
    conn.close()
    if request.method == 'POST':
        allowed, msg = premium_allowed(user, 'crisis')
        if not allowed:
            flash(msg, 'error')
            return redirect(url_for('crisis'))
        topic = request.form.get('topic', '').strip() or 'committee instability'
        country = request.form.get('country', '').strip() or 'Your delegation'
        committee = request.form.get('committee', '').strip() or 'the committee'
        raw_ids = request.form.getlist('file_ids')
        file_ids = [int(x) for x in raw_ids if x.isdigit()]
        context = build_context_text(user['id'], file_ids)
        result = improved_crisis(topic, country, committee, context)
        increment_usage(user['id'], 'usage_crisis')
        log_action(user['id'], 'crisis', f'Generated crisis on {topic}')
    return render_template('crisis.html', result=result, files=files)


@app.route('/debate', methods=['GET', 'POST'])
@login_required
def debate():
    user = current_user()
    result = None
    conn = get_db()
    files = conn.execute('SELECT id, original_name FROM library_files WHERE user_id = ? ORDER BY uploaded_at DESC', (user['id'],)).fetchall()
    conn.close()
    if request.method == 'POST':
        allowed, msg = premium_allowed(user, 'debate')
        if not allowed:
            flash(msg, 'error')
            return redirect(url_for('debate'))
        topic = request.form.get('topic', '').strip() or 'the agenda'
        your_country = request.form.get('your_country', '').strip() or 'Your country'
        opponent_country = request.form.get('opponent_country', '').strip() or 'the opposing delegation'
        raw_ids = request.form.getlist('file_ids')
        file_ids = [int(x) for x in raw_ids if x.isdigit()]
        context = build_context_text(user['id'], file_ids)
        result = improved_debate(topic, your_country, opponent_country, context)
        increment_usage(user['id'], 'usage_debate')
        log_action(user['id'], 'debate', f'Debate against {opponent_country}')
    return render_template('debate.html', result=result, files=files)


@app.route('/admin')
@login_required
@admin_required
def admin_panel():
    conn = get_db()
    admin = current_user()
    users = conn.execute('SELECT * FROM users WHERE id != ? ORDER BY created_at DESC', (admin['id'],)).fetchall()
    platform_stats = {
        'users': conn.execute('SELECT COUNT(*) AS c FROM users').fetchone()['c'],
        'premium': conn.execute("SELECT COUNT(*) AS c FROM users WHERE plan = 'premium'").fetchone()['c'],
        'conferences': conn.execute('SELECT COUNT(*) AS c FROM conferences').fetchone()['c'],
        'files': conn.execute('SELECT COUNT(*) AS c FROM library_files').fetchone()['c'],
    }
    recent_logs = conn.execute('''
        SELECT activity_log.*, users.email
        FROM activity_log
        LEFT JOIN users ON users.id = activity_log.user_id
        ORDER BY activity_log.created_at DESC LIMIT 20
    ''').fetchall()
    conn.close()
    return render_template('admin.html', users=users, platform_stats=platform_stats, recent_logs=recent_logs)


@app.route('/admin/user/<int:user_id>/update', methods=['POST'])
@login_required
@admin_required
def admin_update_user(user_id):
    admin = current_user()
    plan = request.form.get('plan')
    role = request.form.get('role')
    is_active = request.form.get('is_active')
    conn = get_db()
    target = conn.execute('SELECT * FROM users WHERE id = ?', (user_id,)).fetchone()
    if not target:
        conn.close()
        flash('User not found.', 'error')
        return redirect(url_for('admin_panel'))

    if plan in ('free', 'premium'):
        conn.execute('UPDATE users SET plan = ? WHERE id = ?', (plan, user_id))
    if role in ('user', 'admin'):
        conn.execute('UPDATE users SET role = ? WHERE id = ?', (role, user_id))
    if is_active in ('0', '1'):
        conn.execute('UPDATE users SET is_active = ? WHERE id = ?', (int(is_active), user_id))
    conn.commit()
    conn.close()
    log_action(admin['id'], 'admin_update_user', f'Updated user #{user_id}')
    flash('User updated.', 'success')
    return redirect(url_for('admin_panel'))


@app.route('/admin/user/<int:user_id>/delete', methods=['POST'])
@login_required
@admin_required
def admin_delete_user(user_id):
    admin = current_user()
    if user_id == admin['id']:
        flash('You cannot delete your own admin account.', 'error')
        return redirect(url_for('admin_panel'))
    conn = get_db()
    conn.execute('DELETE FROM users WHERE id = ?', (user_id,))
    conn.commit()
    conn.close()
    log_action(admin['id'], 'admin_delete_user', f'Deleted user #{user_id}')
    flash('User deleted.', 'success')
    return redirect(url_for('admin_panel'))


@app.route('/pricing')
@login_required
def pricing():
    return render_template('pricing.html')


if __name__ == '__main__':
    init_db()
    app.run(debug=True)
